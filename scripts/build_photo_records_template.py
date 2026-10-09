"""Build the generic photo-records annex template (docxtpl / Jinja2).

Run once, or whenever the layout changes:
    pixi run --manifest-path ../wd-partner-tasks/src/ecoscope-workflows-ext-wd/pyproject.toml \
        -e test-py312 python scripts/build_photo_records_template.py

Layout follows the Anexo Fotográfico reference (cover page, then one
"Registro" per event: an info table, its photographs two per row, a
"Figura N." caption under each) but carries no partner branding — the
organisation name, title, logo and intro all come from the render
context, so any partner can reuse it, or start from it and restyle.

Context keys (see ecoscope_workflows_ext_wd's generate_photo_records_report):
    organization, title, subtitle, intro, period_label, period, summary,
    generated_on, logo (InlineImage or ""), has_records, no_records,
    labels (dict), records: list of
        {heading, rows: [{label, value}], single, images: [{image, figure, caption}],
         image_rows: [[image, image?], ...]}
"""

from pathlib import Path

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

DST = Path(__file__).parent.parent / "resources" / "templates" / "photo_records_annex_template.docx"

ACCENT = RGBColor(0x2F, 0x5D, 0x50)
MUTED = RGBColor(0x52, 0x60, 0x6D)
LABEL_FILL = "E3EDE8"
BORDER = "C9D6CF"


# ── helpers ───────────────────────────────────────────────────────────────────


def para(container, text: str = "", *, align=None, size=None, bold=False, color=None, italic=False,
         space_before=0, space_after=6, keep_with_next=False):
    p = container.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_with_next
    return p


def tag(container, text: str):
    """A paragraph holding only a docxtpl {%p ... %} tag — removed at render time."""
    p = container.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)
    return p


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table_borders(table, color: str | None):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if color is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_col_widths(table, widths_mm):
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)


def first_par(cell):
    return cell.paragraphs[0]


def add_page_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


def bottom_border(paragraph, color: str):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    p_pr.append(bdr)


def caption(container, figure_expr: str, caption_expr: str):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(figure_expr + " ")
    r.bold = True
    r.font.size = Pt(9)
    r = p.add_run(caption_expr)
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(10)
    return p


def image_par(container, expr: str):
    p = container.add_paragraph(expr)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    return p


# ── document ──────────────────────────────────────────────────────────────────


def build() -> docx.Document:
    doc = docx.Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Mm(20))
    section.top_margin, section.bottom_margin = Mm(22), Mm(20)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("{{ organization }}")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    bottom_border(header, BORDER)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("{{ labels.page }} ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    add_page_field(footer)

    # ── cover ──
    para(doc, space_after=36)
    tag(doc, "{%p if logo %}")
    para(doc, "{{ logo }}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    tag(doc, "{%p endif %}")
    para(doc, "{{ organization }}", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, color=ACCENT, space_after=36)
    para(doc, "{{ title }}", align=WD_ALIGN_PARAGRAPH.CENTER, size=24, bold=True, color=ACCENT, space_after=10)
    para(doc, "{{ subtitle }}", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, space_after=36)

    intro = para(doc, "{{ intro }}", space_after=10)
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    period = doc.add_paragraph()
    period.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = period.add_run("{{ period_label }}: ")
    r.bold = True
    r.font.color.rgb = ACCENT
    period.add_run("{{ period }}")
    period.paragraph_format.space_after = Pt(6)
    para(doc, "{{ summary }}", space_after=6)
    tag(doc, "{%p if not has_records %}")
    para(doc, "{{ no_records }}", italic=True, color=MUTED, space_before=12)
    tag(doc, "{%p endif %}")
    para(doc, "{{ generated_on }}", size=8.5, color=MUTED, space_before=24)

    # ── records ──
    tag(doc, "{%p for r in records %}")
    heading = para(doc, "{{ r.heading }}", size=16, bold=True, color=ACCENT, space_after=8, keep_with_next=True)
    heading.paragraph_format.page_break_before = True

    info = doc.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_borders(info, BORDER)
    first_par(info.cell(0, 0)).text = "{%tr for row in r.rows %}"
    label_cell, value_cell = info.cell(1, 0), info.cell(1, 1)
    run = first_par(label_cell).add_run("{{ row.label }}")
    run.bold = True
    run.font.size = Pt(9.5)
    shade(label_cell, LABEL_FILL)
    first_par(value_cell).add_run("{{ row.value }}").font.size = Pt(9.5)
    first_par(info.cell(2, 0)).text = "{%tr endfor %}"
    set_col_widths(info, (45, 125))

    para(doc, "{{ labels.photographic_record }}", size=12, bold=True, color=ACCENT, space_before=12,
         space_after=4, keep_with_next=True)

    tag(doc, "{%p if r.single %}")
    image_par(doc, "{{ r.images[0].image }}")
    caption(doc, "{{ r.images[0].figure }}", "{{ r.images[0].caption }}")
    tag(doc, "{%p else %}")
    grid = doc.add_table(rows=3, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(grid, None)
    first_par(grid.cell(0, 0)).text = "{%tr for pair in r.image_rows %}"
    # Both cells are built identically (same paragraphs, same spacing) so the
    # two photos start at the same height. generate_photo_records_report
    # always sends 2-item pairs (the 2nd blank for an odd last photo) and
    # sizes both photos in a pair to the same height, so no conditional
    # tags are needed inside the cells.
    tr_pr = grid.rows[1]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))  # keep a pair + captions on one page
    for col in (0, 1):
        cell = grid.cell(1, col)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        img = first_par(cell)
        img.text = f"{{{{ pair[{col}].image }}}}"
        img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.space_before = Pt(4)
        img.paragraph_format.space_after = Pt(2)
        caption(cell, f"{{{{ pair[{col}].figure }}}}", f"{{{{ pair[{col}].caption }}}}")
    first_par(grid.cell(2, 0)).text = "{%tr endfor %}"
    set_col_widths(grid, (85, 85))
    tag(doc, "{%p endif %}")
    tag(doc, "{%p endfor %}")

    return doc


if __name__ == "__main__":
    DST.parent.mkdir(parents=True, exist_ok=True)
    build().save(DST)
    print(f"written: {DST}")
